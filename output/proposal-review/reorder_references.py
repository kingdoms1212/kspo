from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parent
src = root / 'Hi스포링_기획서_AI종합의사결정_수정본.docx'
out = root / 'Hi스포링_기획서_AI종합의사결정_수정본_v2.docx'
d = Document(src)
refs = next(p for p in d.paragraphs if p.text.strip() == '3 활용 데이터명 및 URL')
future = next(p for p in d.paragraphs if p.text.strip() == '4 후속 계획')
body = d.element.body
children = list(body)
block = children[children.index(refs._p):children.index(future._p)]
for element in block:
    body.remove(element)
    body.insert(len(body)-1, element)
def retitle(p, text):
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ''
retitle(future, '3 후속 계획')
retitle(refs, '4 참고자료 및 활용 데이터 출처')
d.save(out)
check = Document(out)
headings = [p.text for p in check.paragraphs if p.style.name == 'Heading 1']
assert headings == ['1 서비스 개요', '2 세부 내용', '3 후속 계획', '4 참고자료 및 활용 데이터 출처'], headings
assert len(check.tables) == len(d.tables) == 3
assert len(check.paragraphs) == len(d.paragraphs)
print('Section order and document structure verified.')
