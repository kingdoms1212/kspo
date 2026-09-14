"""Shared Excel workbook response; feature modules own export columns."""
from io import BytesIO

from django.http import HttpResponse, HttpResponseBadRequest
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Writing the full programme register takes about two minutes and 25MB, which no
# request should hold open. Refusing beats silently truncating a policy document.
EXPORT_ROW_LIMIT = 50000


def over_export_limit(count):
    """Return a refusal when the filtered result is too large to export at once."""
    if count <= EXPORT_ROW_LIMIT:
        return None
    return HttpResponseBadRequest(
        f'조회 결과 {count:,}건은 한 번에 내보낼 수 있는 {EXPORT_ROW_LIMIT:,}건을 넘습니다. '
        '검색 조건을 좁혀 주세요.')


def excel_response(filename, headers, rows):
    """Write filtered rows to XLSX without interpreting source text as formulas.

    Write-only mode keeps cell memory bounded for the full filtered export.
    Text cells preserve leading zeroes in phone numbers and identifiers.
    """
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('조회 결과')
    sheet.freeze_panes = 'A2'
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 40 if header in ('주소', '강좌명', '시설명', '기간') else 22

    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.data_type = 's'
        cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='04386F')
        cell.alignment = Alignment(vertical='center')
        header_cells.append(cell)
    sheet.append(header_cells)
    row_count = 1
    for row in rows:
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # Explicit string type also protects '=...' while preserving the original text.
                cell.data_type = 's'
                cell.number_format = '@'
            cells.append(cell)
        sheet.append(cells)
        row_count += 1
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{row_count}'
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
