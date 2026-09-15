"""Facility controller, register filtering and programme link regressions."""
import csv
import re
import tempfile
from io import BytesIO
from pathlib import Path
from html import unescape
from urllib.parse import parse_qs, quote
from unittest.mock import patch

from openpyxl import load_workbook

from django.test import SimpleTestCase, override_settings
from django.urls import resolve

from .services import facility_flags, facility_owners, facility_states
from . import models, services, transit, views


def facility(**overrides):
    row = dict(id='facility-0', geo_key=('테스트센터', 37.5, 127.0),
               name='테스트센터', flag='공공', state='정상운영', region='서울특별시',
               district='성동구', emd='금호동', address='서울시 1', industry='수영장',
               facility_type='수영장', owner='지자체', operation='자체운영', indoor='실내',
               department='체육과', phone='02-000', area=1200, capacity=None,
               homepage='', national=False)
    row.update(overrides)
    return models.Facility(**row)


REPORT = {'register_rows': 3, 'facilities': 2, 'deleted': 1, 'closed': 0,
          'unlinked': 0, 'source': 'x.csv'}


class FacilityRepositoryTests(SimpleTestCase):
    BASE = ('테스트센터', '공공', '수영장', '수영장', '정상운영',
            '서울특별시 성동구 무수막길 69', '2층', '서울특별시 성동구 금호동2가 1-1',
            '서울특별시', '성동구', '1100000000', '1120000000', '성동구', '금호동2가',
            '자체운영', '지자체', '체육진흥과', '0222040000', '0222047900',
            'https://example.test', '실내', '', '1200', 'N', '37.5518979', '127.0207528', 'N')

    def _load(self, rows):
        models._snapshot.cache_clear()
        self.addCleanup(models._snapshot.cache_clear)
        with tempfile.TemporaryDirectory() as directory:
            with (Path(directory) / models.FACILITY_FILE).open('w', encoding='utf-8-sig', newline='') as handle:
                writer = csv.writer(handle)
                writer.writerow(models.COLUMNS)
                writer.writerows(rows)
            with override_settings(DATA_DIR=directory):
                return models.facilities(), models.load_report()

    def test_deleted_rows_are_excluded_and_counted(self):
        removed = list(self.BASE)
        removed[models.DEL_AT] = 'Y'
        rows, report = self._load([self.BASE, removed])
        self.assertEqual(len(rows), 1)
        self.assertEqual(report, {'register_rows': 2, 'facilities': 1, 'deleted': 1,
                                  'closed': 0, 'unlinked': 0, 'source': models.FACILITY_FILE})

    def test_closed_facilities_are_kept_and_counted_not_dropped(self):
        shut = list(self.BASE)
        shut[models.STATE_VALUE] = '폐업'
        rows, report = self._load([self.BASE, shut])
        self.assertEqual(len(rows), 2)
        self.assertEqual(report['closed'], 1)
        self.assertEqual({row.state for row in rows}, {'정상운영', '폐업'})

    def test_district_comes_from_the_management_name_that_matches_the_address(self):
        rows, _ = self._load([self.BASE])
        self.assertEqual(rows[0].district, '성동구')

    def test_transit_key_is_the_name_and_position(self):
        rows, _ = self._load([self.BASE])
        self.assertEqual(rows[0].geo_key, ('테스트센터', 37.551898, 127.020753))

    def test_missing_road_address_falls_back_without_blocking_the_row(self):
        noaddress = list(self.BASE)
        noaddress[models.RDNMADR_ONE] = ''
        noaddress[models.RDNMADR_TWO] = ''
        rows, _ = self._load([noaddress])
        self.assertEqual(rows[0].address, '서울특별시 성동구 금호동2가 1-1')
        self.assertIsNotNone(rows[0].geo_key)

    def test_a_facility_without_a_position_is_kept_but_cannot_be_linked(self):
        nowhere = list(self.BASE)
        nowhere[models.FCLTY_LA] = '0'
        nowhere[models.FCLTY_LO] = ''
        rows, report = self._load([nowhere])
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0].geo_key)
        self.assertEqual(report['unlinked'], 1)

    def test_address_falls_back_to_administrative_names_as_a_last_resort(self):
        bare = list(self.BASE)
        for index in (models.RDNMADR_ONE, models.RDNMADR_TWO, models.ADDR_ONE):
            bare[index] = ''
        rows, _ = self._load([bare])
        self.assertEqual(rows[0].address, '서울특별시 성동구 금호동2가')

    def test_facility_phone_is_preferred_over_the_department_line(self):
        rows, _ = self._load([self.BASE])
        self.assertEqual(rows[0].phone, '0222047900')
        nophone = list(self.BASE)
        nophone[models.FCLTY_TEL] = ''
        rows, _ = self._load([nophone])
        self.assertEqual(rows[0].phone, '0222040000')

    def test_missing_numbers_stay_missing(self):
        rows, _ = self._load([self.BASE])
        self.assertEqual(rows[0].area, 1200)
        self.assertIsNone(rows[0].capacity)
        self.assertFalse(rows[0].national)
        self.assertEqual(rows[0].flag, '공공')
        self.assertEqual(rows[0].operation, '자체운영')
        self.assertEqual(rows[0].indoor, '실내')


class FacilityScopeFilterTests(SimpleTestCase):
    """Register category and operating state."""

    def setUp(self):
        self.rows = [
            facility(id='facility-0', flag='공공', state='정상운영'),
            facility(id='facility-1', flag='신고', state='정상운영'),
            facility(id='facility-2', flag='신고', state='폐업'),
        ]

    def _get(self, query):
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            return self.client.get('/facilities', query)

    def test_options_come_from_the_register(self):
        self.assertEqual(facility_flags(self.rows), ['공공', '신고'])
        self.assertEqual(facility_states(self.rows), ['정상운영', '폐업'])

    def test_category_and_state_narrow_the_list(self):
        self.assertEqual([r.id for r in services.filter_facilities(self.rows, flag='공공')], ['facility-0'])
        self.assertEqual([r.id for r in services.filter_facilities(self.rows, state='폐업')], ['facility-2'])
        self.assertEqual([r.id for r in services.filter_facilities(self.rows, flag='신고', state='정상운영')],
                         ['facility-1'])

    def test_closed_facilities_are_hidden_by_default_and_the_select_says_so(self):
        """A default that narrows must be visible, not silent."""
        page = self._get({})
        self.assertEqual([r.id for r in page.context['rows']], ['facility-0', 'facility-1'])
        self.assertEqual(page.context['params']['state'], '정상운영')
        self.assertContains(page, '<option selected>정상운영</option>', html=True)

    def test_an_explicit_blank_state_widens_to_every_row(self):
        page = self._get({'state': ''})
        self.assertEqual([r.id for r in page.context['rows']], ['facility-0', 'facility-1', 'facility-2'])
        self.assertEqual(page.context['params']['state'], '')

    def test_export_follows_the_same_default(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            default = self.client.get('/export/facilities.xlsx')
            widened = self.client.get('/export/facilities.xlsx', {'state': ''})
        self.assertEqual(len(list(load_workbook(BytesIO(default.content)).active.values)), 3)
        self.assertEqual(len(list(load_workbook(BytesIO(widened.content)).active.values)), 4)


class FacilityOwnerFilterTests(SimpleTestCase):
    def setUp(self):
        self.rows = [
            facility(id='facility-0', owner='지자체'),
            facility(id='facility-1', owner='국민체육'),
            facility(id='facility-2', owner='보유주체 미제공'),
        ]

    def test_options_skip_the_missing_value_placeholder(self):
        self.assertEqual(facility_owners(self.rows), ['국민체육', '지자체'])

    def test_filter_narrows_to_one_owning_body(self):
        found = services.filter_facilities(self.rows, owner='국민체육')
        self.assertEqual([row.id for row in found], ['facility-1'])

    def test_page_and_export_apply_the_same_owner(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'owner': '국민체육'})
            export = self.client.get('/export/facilities.xlsx', {'owner': '국민체육'})
        self.assertEqual([row.id for row in page.context['rows']], ['facility-1'])
        self.assertContains(page, '<option value="">전체 보유주체</option>', html=True)
        records = list(load_workbook(BytesIO(export.content)).active.values)
        self.assertEqual(len(records), 2)
        self.assertEqual(records[1][10], '국민체육')

    def test_list_shows_the_owner_even_for_a_national_centre(self):
        """The national-centre flag is a separate column and must not replace the owner.

        147 facilities carry the flag, and they are almost exactly the ones the
        owner filter exists to surface, so hiding the owner there made the
        filter disagree with the column it filtered on.
        """
        rows = [facility(id='facility-0', owner='법무부', national=True),
                facility(id='facility-1', owner='대한체육회', national=True)]
        with patch('app.facilities.views.models.facilities', return_value=rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'owner': '법무부'})
        self.assertEqual([row.id for row in page.context['rows']], ['facility-0'])
        body = page.content.decode()
        self.assertIn('법무부<small>국민체육센터</small>', body)
        # The other owner stays in the select options but leaves the table.
        self.assertNotIn('<td>대한체육회', body)
        self.assertIn('<option >대한체육회</option>', body)

    def test_detail_uses_the_shared_photo_with_a_caption_that_says_so(self):
        """One stand-in image serves every facility, so it must not read as this one's."""
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'facilityId': 'facility-0'})
        self.assertContains(page, 'src="/static/image/center.jpg"')
        self.assertContains(page, '대표 이미지')
        self.assertContains(page, 'alt=""')

    def test_facility_link_sets_facility_id_exactly_once(self):
        rows = [facility(id='facility-0'), facility(id='facility-1')]
        with patch('app.facilities.views.models.facilities', return_value=rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'owner': '지자체', 'facilityId': 'facility-0'})
        link = re.search(r'href="(\?[^"]*facilityId=facility-1[^"]*)"', page.content.decode()).group(1)
        parsed = parse_qs(unescape(link).lstrip('?'))
        self.assertEqual(parsed['facilityId'], ['facility-1'])
        self.assertEqual(parsed['owner'], ['지자체'])

    def test_owner_survives_paging_and_facility_selection(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'owner': '지자체', 'facilityId': 'facility-0'})
        self.assertIn('owner=' + quote('지자체'), page.context['page_query'])
        self.assertEqual(page.context['selected'].id, 'facility-0')


class FacilityTransitTests(SimpleTestCase):
    """Nearby stops are matched on position, never on an address or a name."""

    HEADERS = transit.COLUMNS
    BASE = ('테스트센터', '37.5518979', '127.0207528', '버스', '120.0', '150.0', '180', '금호역앞')

    def setUp(self):
        services.cache_clear()
        self.addCleanup(services.cache_clear)

    def _load(self, rows, facility_row=None):
        facility_row = facility_row or facility()
        with tempfile.TemporaryDirectory() as directory:
            with (Path(directory) / transit.TRANSIT_FILE).open('w', encoding='utf-8-sig', newline='') as handle:
                writer = csv.writer(handle)
                writer.writerow(self.HEADERS)
                writer.writerows(rows)
            with override_settings(DATA_DIR=directory),                  patch('app.facilities.services.models.facilities', return_value=[facility_row]):
                return services.facility_transit(facility_row)

    def _facility(self, **kw):
        return facility(geo_key=transit.geo_key('테스트센터', '37.5518979', '127.0207528'), **kw)

    def test_position_rounds_to_six_decimals_and_rejects_a_missing_one(self):
        self.assertEqual(transit.geo_key(' 테스트 센터 ', '37.55189794', '127.02075281'),
                         ('테스트 센터', 37.551898, 127.020753))
        self.assertIsNone(transit.geo_key('', '37.5', '127.0'))
        self.assertIsNone(transit.geo_key('이름', '0', '127.0'))
        self.assertIsNone(transit.geo_key('이름', '', '127.0'))
        self.assertIsNone(transit.geo_key('이름', 'x', '127.0'))

    def test_stops_are_ordered_by_straight_distance_and_counted_by_mode(self):
        far = list(self.BASE); far[4], far[6], far[7] = '900.0', '600', '먼정류장'
        subway = list(self.BASE); subway[3], subway[4], subway[6], subway[7] = '지하철', '400.0', '300', '금호역'
        result = self._load([far, self.BASE, subway], self._facility())
        self.assertEqual([s['name'] for s in result['stops']], ['금호역앞', '금호역', '먼정류장'])
        self.assertEqual((result['bus'], result['subway'], result['total']), (2, 1, 3))
        self.assertEqual(result['nearest_metres'], 120)
        self.assertEqual(result['nearest_minutes'], 3)

    def test_a_close_stop_outranks_one_that_merely_records_a_walk_time(self):
        """53.9% of rows carry no walk time, so distance is the only shared order."""
        close = list(self.BASE); close[4], close[5], close[6], close[7] = '20.0', '', '', '바로앞'
        walked = list(self.BASE); walked[4], walked[6], walked[7] = '800.0', '900', '먼정류장'
        result = self._load([walked, close], self._facility())
        self.assertEqual([s['name'] for s in result['stops']], ['바로앞', '먼정류장'])
        self.assertIsNone(result['stops'][0]['minutes'])
        self.assertEqual(result['walk_known'], 1)

    def test_another_site_with_the_same_name_is_not_borrowed(self):
        elsewhere = list(self.BASE)
        elsewhere[1], elsewhere[2], elsewhere[7] = '35.1', '129.0', '남의정류장'
        result = self._load([elsewhere], self._facility())
        self.assertFalse(result['has_records'])
        self.assertEqual(result['stops'], ())

    def test_a_facility_without_a_position_reports_without_reading_the_file(self):
        with patch('app.facilities.transit._index') as index,              patch('app.facilities.services.models.facilities', return_value=[]):
            result = services.facility_transit(facility(geo_key=None))
        index.assert_not_called()
        self.assertEqual(result['error'], transit.NO_POSITION)
        self.assertFalse(result['has_records'])

    def test_a_missing_source_file_reports_instead_of_raising(self):
        with tempfile.TemporaryDirectory() as directory:
            with override_settings(DATA_DIR=directory),                  patch('app.facilities.services.models.facilities', return_value=[self._facility()]):
                result = services.facility_transit(self._facility())
        self.assertEqual(result['error'], transit.MISSING_SOURCE)

    def test_only_the_nearest_are_kept_and_the_total_still_reported(self):
        rows = []
        for index in range(transit.KEEP + 4):
            row = list(self.BASE)
            row[4], row[7] = f'{100 + index}.0', f'정류장{index:02}'
            rows.append(row)
        result = self._load(rows, self._facility())
        self.assertEqual(result['total'], transit.KEEP + 4)
        self.assertEqual(result['shown'], transit.KEEP)
        self.assertEqual(result['stops'][0]['name'], '정류장00')

    def test_a_missing_walk_time_stays_missing_rather_than_being_estimated(self):
        blank = list(self.BASE); blank[4], blank[5], blank[6], blank[7] = '300.0', '', '', '시간미상'
        result = self._load([blank, self.BASE], self._facility())
        self.assertEqual([s['name'] for s in result['stops']], ['금호역앞', '시간미상'])
        self.assertIsNone(result['stops'][1]['minutes'])
        self.assertEqual(result['stops'][1]['metres'], 300)
        self.assertEqual(result['nearest_minutes'], 3)


class FacilityViewTests(SimpleTestCase):
    def setUp(self):
        self.rows = [facility(id='facility-0', name='=1+1', region='서울특별시'),
                     facility(id='facility-1', name='부산센터', region='부산광역시', industry='체육관')]

    def test_module_routes_and_export_escaping(self):
        self.assertIs(resolve('/facilities').func, views.facilities)
        self.assertIs(resolve('/export/facilities.xlsx').func, views.export_facilities)
        self.assertIs(resolve('/export/facility-transit.xlsx').func, views.export_facility_transit)
        with patch('app.facilities.views.models.facilities', return_value=self.rows), \
             patch('app.facilities.views.models.load_report', return_value=REPORT), \
             patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'region': '서울특별시', 'facilityId': 'facility-0'})
            export = self.client.get('/export/facilities.xlsx', {'region': '서울특별시'})
            empty = self.client.get('/export/facilities.xlsx', {'region': '대구광역시'})
        self.assertTemplateUsed(page, 'facilities/index.html')
        self.assertEqual(page.context['selected'].name, '=1+1')
        self.assertContains(page, '=1+1')
        workbook = load_workbook(BytesIO(export.content))
        records = list(workbook.active.values)
        self.assertEqual(records[1][1], '=1+1')
        self.assertEqual(workbook.active['B2'].data_type, 's')
        self.assertEqual(workbook.active.freeze_panes, 'A2')
        self.assertEqual(workbook.active.auto_filter.ref, 'A1:T2')
        self.assertIn('SPORT_INSIGHT_facilities.xlsx', export['Content-Disposition'])
        self.assertEqual(len(list(load_workbook(BytesIO(empty.content)).active.values)), 1)

    def test_export_button_is_disabled_with_a_reason_over_the_limit(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows),              patch('app.facilities.views.models.load_report', return_value=REPORT),              patch('app.facilities.views.facility_transit', return_value=None),              patch('app.common.exports.EXPORT_ROW_LIMIT', 1),              patch('app.facilities.views.EXPORT_ROW_LIMIT', 1):
            page = self.client.get('/facilities')
        self.assertContains(page, 'class="button export" type="button" disabled')
        self.assertContains(page, 'id="facilities-export-limit"')
        self.assertContains(page, '한 번에 내보낼 수 있는 1건을 넘습니다')
        self.assertNotContains(page, 'href="/export/facilities.xlsx')

    def test_industry_filter_scopes_page_and_export_alike(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows), \
             patch('app.facilities.views.models.load_report', return_value=REPORT), \
             patch('app.facilities.views.facility_transit', return_value=None):
            page = self.client.get('/facilities', {'industry': '체육관'})
            export = self.client.get('/export/facilities.xlsx', {'industry': '체육관'})
        self.assertEqual([row.id for row in page.context['rows']], ['facility-1'])
        self.assertEqual(len(list(load_workbook(BytesIO(export.content)).active.values)), 2)

    def test_unselected_and_filtered_out_facilities_do_not_load_transit(self):
        with patch('app.facilities.views.models.facilities', return_value=self.rows), \
             patch('app.facilities.views.models.load_report', return_value=REPORT), \
             patch('app.facilities.services.transit.stops_for') as lookup:
            for query in ({}, {'facilityId': 'missing'}, {'facilityId': 'facility-0', 'query': '없는시설'}):
                response = self.client.get('/facilities', query)
                self.assertIsNone(response.context['selected'])
                self.assertIsNone(response.context['transit'])
            # The 1.6M-row transit file is never touched without a selection.
            lookup.assert_not_called()
        self.assertContains(response, '시설을 찾을 수 없습니다')
