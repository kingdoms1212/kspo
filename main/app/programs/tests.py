"""Program module regressions. CSV access is mocked at the repository boundary."""
import csv
import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from django.test import SimpleTestCase, override_settings
from django.urls import resolve, reverse

from . import models, views
from .services import (deduplicate_programs, filter_programs,
                       program_district_distribution, program_facility_types,
                       selected_weekdays)


def program(**overrides):
    row = dict(id='program-0', facility='시설', region='서울', district='중구', address='주소',
               phone='02-000', facility_type='수영장', name='A', sport='수영', target='성인',
               weekday='월수금', time='19:00~19:50', period='2025-03-01 ~ 2025-03-31',
               capacity=20, fee=50000, fee_unit='단위 미확인', homepage='', walk_minutes=3, stop='정류장')
    row.update(overrides)
    return models.Program(**row)


REPORT = {'register_rows': 5, 'programs': 3, 'duplicates': 2,
          'facilities': 2, 'source': 'x.csv'}


class ProgramServiceTests(SimpleTestCase):
    def setUp(self):
        self.rows = [program(id='b', name='B', region='서울', fee=2000, walk_minutes=9,
                             capacity=5, target='청소년', weekday='화목', facility_type='체육관'),
                     program(id='a', name='A', region='부산', fee=1000, walk_minutes=1,
                             capacity=50, target='성인', weekday='월수금', facility_type='수영장')]

    def test_sort_does_not_mutate_cached_repository_rows(self):
        with patch('app.programs.services.models.programs', return_value=self.rows):
            result = filter_programs({'sort': 'fee'})
        self.assertEqual([row.id for row in result], ['a', 'b'])
        self.assertEqual([row.id for row in self.rows], ['b', 'a'])
        self.assertIsNot(result, self.rows)

    def test_walk_and_capacity_sorts_keep_missing_values_last(self):
        rows = [*self.rows, program(id='c', name='C', walk_minutes=None, capacity=None)]
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual([r.id for r in filter_programs({'sort': 'walk'})], ['a', 'b', 'c'])
            self.assertEqual([r.id for r in filter_programs({'sort': 'capacity'})], ['a', 'b', 'c'])

    def test_program_and_facility_name_filters_are_independent(self):
        rows = [
            program(id='a', name='아침 수영', facility='시민 체육관'),
            program(id='b', name='저녁 수영', facility='가족 수영장'),
            program(id='c', name='아침 요가', facility='시민 문화관'),
        ]
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual(
                [row.id for row in filter_programs({'program_query': '아침'})],
                ['a', 'c'],
            )
            self.assertEqual(
                [row.id for row in filter_programs({'facility_query': '시민'})],
                ['a', 'c'],
            )
            self.assertEqual(
                [row.id for row in filter_programs({
                    'program_query': '수영',
                    'facility_query': '가족',
                })],
                ['b'],
            )

    def test_new_sort_options_order_by_the_requested_field(self):
        rows = [
            program(id='a', name='Z', facility='B', region='C', district='A', sport='C', fee=3000),
            program(id='b', name='A', facility='C', region='A', district='A', sport='A', fee=1000),
            program(id='c', name='M', facility='A', region='B', district='A', sport='B', fee=2000),
        ]
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual([row.id for row in filter_programs({'sort': 'name'})], ['b', 'c', 'a'])
            self.assertEqual([row.id for row in filter_programs({'sort': 'facility'})], ['c', 'a', 'b'])
            self.assertEqual([row.id for row in filter_programs({'sort': 'region'})], ['b', 'c', 'a'])
            self.assertEqual([row.id for row in filter_programs({'sort': 'sport'})], ['b', 'c', 'a'])
            self.assertEqual([row.id for row in filter_programs({'sort': 'fee'})], ['b', 'c', 'a'])
            self.assertEqual(
                [row.id for row in filter_programs({'sort': 'name_desc'})],
                ['a', 'c', 'b'],
            )
            self.assertEqual(
                [row.id for row in filter_programs({'sort': 'fee_desc'})],
                ['a', 'c', 'b'],
            )

    def test_target_weekday_and_facility_type_filters(self):
        with patch('app.programs.services.models.programs', return_value=self.rows):
            self.assertEqual([r.id for r in filter_programs({'target': '청소년'})], ['b'])
            self.assertEqual([r.id for r in filter_programs({'weekday': '월수금'})], ['a'])
            self.assertEqual([r.id for r in filter_programs({'facility_type': '체육관'})], ['b'])
        self.assertEqual(program_facility_types(self.rows), ['수영장', '체육관'])

    def test_multiple_weekday_filter_requires_an_exact_weekday_match(self):
        rows = [
            program(id='a', weekday='월금'),
            program(id='b', weekday='월화'),
            program(id='c', weekday='수금'),
            program(id='d', weekday='월화수목금토일'),
            program(id='e', weekday='금, 월'),
        ]
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual(
                [row.id for row in filter_programs({'weekday': '금월'})],
                ['a', 'e'],
            )
        self.assertEqual(selected_weekdays('금월월'), ('월', '금'))

    def test_single_weekday_filter_matches_programs_containing_that_day(self):
        rows = [
            program(id='a', weekday='월'),
            program(id='b', weekday='월수금'),
            program(id='c', weekday='화목'),
        ]
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual(
                [row.id for row in filter_programs({'weekday': '월'})],
                ['a', 'b'],
            )

    def test_visible_duplicate_programs_keep_only_the_first_sorted_row(self):
        rows = [
            program(id='b', period='2026-08-01 ~ 2026-08-31', time='20:00~20:50'),
            program(id='a', period='2026-07-01 ~ 2026-07-31', time='19:00~19:50'),
            program(id='c', fee=60000),
        ]

        self.assertEqual([row.id for row in deduplicate_programs(rows)], ['b', 'c'])
        with patch('app.programs.services.models.programs', return_value=rows):
            self.assertEqual([row.id for row in filter_programs({'sort': 'name'})], ['a', 'c'])

    def test_district_distribution_groups_every_region(self):
        rows = [
            program(region='서울특별시', district='강남구'),
            program(region='부산광역시', district='해운대구'),
            program(region='부산광역시', district='해운대구'),
            program(region='경기도', district='수원시 영통구'),
        ]
        self.assertEqual(program_district_distribution(rows), {
            '경기도': [('수원시 영통구', 1)],
            '부산광역시': [('해운대구', 2)],
            '서울특별시': [('강남구', 1)],
        })


class ProgramRepositoryTests(SimpleTestCase):
    BASE = ('1100000000', '서울특별시', '1120000000', '성동구', '테스트센터',
            '서울특별시 성동구 무수막길 69', '0222047900', '수영장', '수영', '저녁수영',
            '성인,<br>청소년(13세이상)', '20260701', '20260731', '월화수목금', '19:00-19:50',
            '25', '49500.00000', '', 'https://example.test', '금호역', '600', '120', '', '', '')

    def _load(self, rows):
        models._snapshot.cache_clear()
        self.addCleanup(models._snapshot.cache_clear)
        with tempfile.TemporaryDirectory() as directory:
            with (Path(directory) / models.PROGRAM_FILE).open('w', encoding='utf-8-sig', newline='') as handle:
                writer = csv.writer(handle)
                writer.writerow(models.COLUMNS)
                writer.writerows(rows)
            with override_settings(DATA_DIR=directory):
                return models.programs(), models.load_report()

    def test_identical_rows_collapse_and_are_counted(self):
        different = list(self.BASE)
        different[models.TIZN] = '20:00-20:50'
        rows, report = self._load([self.BASE, self.BASE, different])
        self.assertEqual(len(rows), 2)
        self.assertEqual(report['duplicates'], 1)
        self.assertEqual(report['register_rows'], 3)
        self.assertEqual(report['facilities'], 1)

    def test_source_fields_are_normalised_without_inventing_values(self):
        rows, _ = self._load([self.BASE])
        row = rows[0]
        self.assertEqual(row.target, '성인, 청소년(13세이상)')
        self.assertEqual(row.period, '2026-07-01 ~ 2026-07-31')
        self.assertEqual(row.fee, 49500)
        self.assertEqual(row.fee_unit, '단위 미확인')
        self.assertEqual(row.capacity, 25)
        self.assertEqual(row.walk_minutes, 2)

    def test_blank_programme_type_falls_back_to_the_facility_industry(self):
        blank = list(self.BASE)
        blank[models.PROGRM_TY_NM] = ''
        rows, _ = self._load([blank])
        self.assertEqual(rows[0].sport, '수영장')

    def test_missing_walk_times_stay_missing(self):
        nowalk = list(self.BASE)
        for index in models.WALK:
            nowalk[index] = ''
        rows, _ = self._load([nowalk])
        self.assertIsNone(rows[0].walk_minutes)


class ProgramViewTests(SimpleTestCase):
    def setUp(self):
        self.rows = [program(id='b', name='B', region='서울', sport='축구', weekday='화',
                             fee=1000, target='유아'),
                     program(id='a', name='A', region='부산', sport='수영', weekday='월',
                             fee=2000, target='유아')]

    @override_settings(PROGRAM_REGION_FILTER_MODE='selectable')
    def test_page_and_export_use_same_filters(self):
        with patch('app.programs.models.programs', return_value=self.rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            response = self.client.get(reverse('programs'), {'region': '서울'})
            export = self.client.get(reverse('export_programs'), {'region': '서울'})
        self.assertTemplateUsed(response, 'programs/index.html')
        self.assertEqual(response.context['result_count'], 1)
        workbook = load_workbook(BytesIO(export.content))
        records = list(workbook.active.values)
        self.assertEqual(export['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('SPORT_INSIGHT_programs.xlsx', export['Content-Disposition'])
        self.assertContains(response, '엑셀 내보내기')
        self.assertEqual(len(records), 2)
        self.assertEqual(records[1][0], response.context['page_obj'][0].id)
        self.assertIn('정류장 도보(분)', records[0])

    def test_page_and_export_share_visible_field_deduplication(self):
        rows = [
            program(id='b', period='2026-08-01 ~ 2026-08-31', time='20:00~20:50'),
            program(id='a', period='2026-07-01 ~ 2026-07-31', time='19:00~19:50'),
        ]
        with patch('app.programs.models.programs', return_value=rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            response = self.client.get(reverse('programs'))
            export = self.client.get(reverse('export_programs'))

        self.assertEqual(response.context['result_count'], 1)
        self.assertEqual([row.id for row in response.context['page_obj']], ['a'])
        records = list(load_workbook(BytesIO(export.content)).active.values)
        self.assertEqual(len(records), 2)
        self.assertEqual(records[1][0], 'a')

    def test_page_states_its_own_source_and_deduplication(self):
        with patch('app.programs.models.programs', return_value=self.rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            response = self.client.get(reverse('programs'))
        # The page names the source it actually loaded, from the load report.
        self.assertContains(response, '출처: x.csv')
        self.assertContains(response, '중복 2행을 제외해 3건을 보존했습니다')
        self.assertContains(response, '프로그램명, 시설, 지역, 종목, 요일, 수강료가 모두 같은 강좌를 1건으로 표시합니다')

    def test_pagination_uses_full_navigation_and_returns_to_comparison(self):
        rows = [program(id=str(index), name=f'강좌 {index:03}') for index in range(221)]
        with patch('app.programs.models.programs', return_value=rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            response = self.client.get(reverse('programs'))

        body = response.content.decode()
        pagination = body.split('aria-label="프로그램 비교표 페이지"', 1)[1].split('</nav>', 1)[0]
        self.assertIn('page=2#program-comparison', pagination)
        self.assertIn('page=11#program-comparison', pagination)
        self.assertIn('aria-label="10페이지 이전"', pagination)
        self.assertIn('aria-label="10페이지 다음"', pagination)
        self.assertIn('&lt;&lt;', pagination)
        self.assertIn('&gt;&gt;', pagination)
        self.assertNotIn('hx-get=', pagination)
        self.assertNotIn('hx-target=', pagination)

        with patch('app.programs.models.programs', return_value=rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            page_eleven = self.client.get(reverse('programs'), {'page': 11})
        page_eleven_pagination = page_eleven.content.decode().split(
            'aria-label="프로그램 비교표 페이지"', 1
        )[1].split('</nav>', 1)[0]
        self.assertIn('page=1#program-comparison', page_eleven_pagination)

    @override_settings(PROGRAM_REGION_FILTER_MODE='selectable')
    def test_oversized_export_is_refused_instead_of_truncated(self):
        rows = [program(id=str(i), name=f'강좌 {i}') for i in range(3)]
        with patch('app.programs.models.programs', return_value=rows),              patch('app.programs.models.load_report', return_value=REPORT),              patch('app.common.exports.EXPORT_ROW_LIMIT', 2),              patch('app.programs.views.EXPORT_ROW_LIMIT', 2):
            page = self.client.get(reverse('programs'))
            refused = self.client.get(reverse('export_programs'))
            allowed = self.client.get(reverse('export_programs'), {'region': '부산'})
        self.assertEqual(refused.status_code, 400)
        self.assertIn('검색 조건을 좁혀', refused.content.decode())
        self.assertEqual(allowed.status_code, 200)
        # The button stays visible but disabled, with the reason on a focusable
        # help control beside it: a disabled button cannot hold a tooltip.
        self.assertContains(page, 'class="button export" type="button" disabled')
        self.assertContains(page, 'id="programs-export-limit"')
        self.assertContains(page, '한 번에 내보낼 수 있는 2건을 넘습니다')
        self.assertContains(page, 'aria-describedby="programs-export-limit"')
        self.assertNotContains(page, 'href="/export/programs.xlsx')

    def test_export_button_is_a_plain_link_while_under_the_limit(self):
        with patch('app.programs.models.programs', return_value=self.rows),              patch('app.programs.models.load_report', return_value=REPORT):
            page = self.client.get(reverse('programs'))
        self.assertContains(page, 'href="/export/programs.xlsx')
        self.assertNotContains(page, 'type="button" disabled')
        self.assertNotContains(page, 'id="programs-export-limit"')

    @override_settings(PROGRAM_REGION_FILTER_MODE='selectable')
    def test_program_page_exposes_dependent_district_options(self):
        rows = [
            program(id='1', name='수영', region='서울특별시', district='관악구', sport='수영'),
            program(id='2', name='축구', region='서울특별시', district='금천구', sport='축구'),
            program(id='3', name='농구', region='경기도', district='수원시', sport='농구'),
        ]
        with patch('app.programs.models.programs', return_value=rows):
            response = self.client.get(reverse('programs'), {
                'region': '서울특별시',
                'district': '관악구',
            })

        self.assertEqual(response.context['districts'], ['관악구', '금천구'])
        self.assertEqual(response.context['params']['district'], '관악구')
        self.assertEqual(response.context['region_district_map']['경기도'], ['수원시'])
        self.assertContains(response, 'id="district"')
        self.assertContains(response, 'id="region-district-map"')

    @override_settings(PROGRAM_REGION_FILTER_MODE='fixed')
    def test_fixed_region_mode_hides_region_and_keeps_all_districts(self):
        rows = [
            program(id='1', region='서울특별시', district='관악구'),
            program(id='2', region='서울특별시', district='금천구'),
        ]
        with patch('app.programs.models.programs', return_value=rows):
            response = self.client.get(reverse('programs'), {
                'region': '부산광역시',
                'district': '관악구',
            })

        self.assertNotContains(response, '<label for="region">지역</label>')
        self.assertContains(response, '<label for="district">지역</label>')
        self.assertContains(response, '<option value="">전체</option>')
        self.assertEqual(response.context['params']['region'], '')
        self.assertEqual(response.context['districts'], ['관악구', '금천구'])
        self.assertEqual(response.context['result_count'], 1)

    def test_weekday_picker_includes_monday_and_updates_the_weekday_field(self):
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'), {'weekday': '월금'})

        body = response.content.decode()
        picker = body.split('class="weekday-options"', 1)[1].split('</div>', 1)[0]
        for day in ('월', '화', '수', '목', '금', '토', '일'):
            self.assertIn(f'value="{day}"', picker)
            self.assertIn(f'<span>{day}</span>', picker)
        self.assertContains(response, 'id="weekday" name="weekday" type="hidden" value="월금"')
        self.assertContains(response, 'id="weekday-label-value">월, 금</span>')
        self.assertIn('weekdayInput.value = selectedDays.join(\'\');', body)

    def test_program_page_preserves_budget_range_without_removed_fields(self):
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'), {
                'budget_min': '1000',
                'budget_max': '2000',
            })

        self.assertTrue(response.context['budget_plan']['active'])
        self.assertContains(response, 'name="budget_min"')
        self.assertContains(response, 'value="1000"')
        self.assertNotContains(response, 'name="planned_people"')
        self.assertNotContains(response, 'name="support_months"')
        self.assertNotContains(response, 'name="target"')
        self.assertNotContains(response, 'class="budget-filter"')
        self.assertNotContains(response, '조건을 입력한 후 조회를 눌러 적용하세요')
        self.assertNotContains(response, '수강료 단위는 별도 확인이 필요합니다')

    def test_program_page_declares_the_shared_region_map(self):
        """The page declares a map; static/region-map.js owns the drawing."""
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'))

        self.assertContains(response, '지역 분포')
        self.assertContains(response, 'data-region-map')
        self.assertContains(response, 'id="program-region-map"')
        self.assertContains(response, 'id="program-region-map-region-data"')
        self.assertContains(response, 'id="program-region-map-district-data"')
        self.assertContains(response, 'id="program-region-map-back"')
        self.assertContains(response, 'data-region-map-back-suppressed')
        self.assertContains(response, 'data-region-click="drilldown"')
        self.assertContains(response, 'region-map.js')
        self.assertContains(response, 'programs-region-map.js')
        self.assertContains(response, 'id="program-scroll-position"')
        self.assertContains(response, 'aria-describedby="program-distribution-heading-tooltip"')
        self.assertContains(response, '시도를 선택하면 해당 지역의 시군구별 프로그램 분포를 확인할 수 있습니다.')
        self.assertContains(response, '지도 색상은 검색된 프로그램의 지역별 건수이며 시설 중복을 제거한 수치는 아닙니다.')
        self.assertContains(response, 'id="program-region-map-help" class="region-map-help sr-only"')
        # The drawing details moved out of the page with the library.
        self.assertNotContains(response, 'echarts@5.6.0')
        self.assertNotContains(response, "backgroundColor: '#e6edf5'")


    def test_program_heading_description_uses_accessible_tooltip(self):
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'))

        self.assertContains(response, 'class="heading-help"')
        self.assertContains(response, 'aria-describedby="program-heading-tooltip"')
        self.assertContains(response, 'id="program-heading-tooltip" class="heading-tooltip" role="tooltip"')
        self.assertContains(response, '서울의 시군구와 종목별로 등록 강좌를 조회하고 비교합니다.')
        self.assertContains(response, 'href="#i-alert"')

    def test_comparison_source_note_uses_accessible_tooltip(self):
        with patch('app.programs.models.programs', return_value=self.rows), \
             patch('app.programs.models.load_report', return_value=REPORT):
            response = self.client.get(reverse('programs'))

        self.assertContains(response, 'aria-describedby="program-comparison-heading-tooltip"')
        self.assertContains(response, 'id="program-comparison-heading-tooltip" class="heading-tooltip" role="tooltip"')
        self.assertContains(response, '출처: x.csv, 원본 5행 중 모든 적재 항목이 같은 중복 2행을 제외해 3건을 보존했습니다.')
        self.assertContains(response, '내보내기는 모든 페이지의 적용 결과입니다.')

    def test_distribution_replaces_top_program_cards(self):
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'))

        body = response.content.decode()
        self.assertNotIn('조건에 맞는 프로그램', body)
        self.assertNotIn('class="program-cards"', body)
        self.assertLess(body.index('지역 분포'), body.index('프로그램 비교표'))

    def test_comparison_table_hides_target_and_formats_fee(self):
        with patch('app.programs.models.programs', return_value=self.rows):
            response = self.client.get(reverse('programs'))

        comparison_table = response.content.decode().split('aria-label="프로그램 비교표 가로 스크롤"', 1)[1].split('</table>', 1)[0]
        self.assertIn('<th scope="col">종목</th>', comparison_table)
        self.assertIn('<th scope="col">수강료</th>', comparison_table)
        self.assertIn('<td class="weekday-cell">월</td>', comparison_table)
        self.assertIn('<td class="weekday-cell">화</td>', comparison_table)
        self.assertIn('<td class="price">1,000원</td>', comparison_table)
        self.assertIn('<td class="price">2,000원</td>', comparison_table)
        self.assertNotIn('종목 / 대상', comparison_table)
        self.assertNotIn('수강료 / 단위', comparison_table)
        self.assertNotIn('청소년', comparison_table)
        self.assertNotIn('성인', comparison_table)
        self.assertNotIn('단위 미확인', comparison_table)

    def test_url_resolves_to_module_controller(self):
        self.assertIs(resolve('/programs').func, views.programs)
        self.assertIs(resolve('/export/programs.xlsx').func, views.export_programs)
